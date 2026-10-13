"""PPT 解析器：逐页（幻灯片）提取文本框/表格文字，保留页码

两层策略（业界常规 OCR fallback）：
1. 文字层：python-pptx 读文本框/表格（正常 PPT 零损耗）；
2. OCR fallback：整份文字层为空 → 判定为"图片型 PPT"（整页贴图，
   老师课件常见形态），直接取出幻灯片内嵌的图片逐张 OCR。

为什么直接取内嵌图片而不是渲染幻灯片：
- 渲染 PPT 需要 PowerPoint/LibreOffice（重依赖）；实测 PyMuPDF 虽能打开
  pptx 但渲染不出图形内容（空白页）；
- 图片型 PPT 的内容全在内嵌图里，取图 OCR 等价且零新依赖。
"""

from pathlib import Path

from pptx import Presentation

from wisnote.parsers.base import Block

MAX_OCR_SLIDES = 200


def _extract_text_layer(prs) -> list[list[str]]:
    """每页的文本框/表格文字列表"""
    pages: list[list[str]] = []
    for slide in prs.slides:
        texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    t = "".join(run.text for run in p.runs).strip()
                    if t:
                        texts.append(t)
            elif getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        t = cell.text.strip()
                        if t:
                            texts.append(t)
        pages.append(texts)
    return pages


def _picture_blobs(slide) -> list[bytes]:
    """幻灯片内嵌图片的二进制列表（跳过读不出的损坏图）"""
    blobs = []
    for shape in slide.shapes:
        if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
            try:
                blobs.append(shape.image.blob)
            except Exception:  # nosec B112 单张损坏图跳过，不影响其余图片的 OCR
                continue
    return blobs


def parse_pptx(path: Path) -> list[Block]:
    blocks: list[Block] = []
    para = 0
    prs = Presentation(str(path))
    pages = _extract_text_layer(prs)

    # OCR fallback：整份文字层为空 → 图片型 PPT，取内嵌图逐张 OCR
    if not any(pages):
        if len(pages) > MAX_OCR_SLIDES:
            raise ValueError(
                f"图片型 PPT 共 {len(pages)} 页，超过 OCR 上限 {MAX_OCR_SLIDES} 页，"
                "请拆分后再上传"
            )
        from wisnote.parsers.image_parser import ocr_image_bytes

        for slide_no, slide in enumerate(prs.slides, start=1):
            for blob in _picture_blobs(slide):
                for b in ocr_image_bytes(blob, page=slide_no):
                    para += 1
                    blocks.append(Block(page=slide_no, para_index=para, text=b.text))
        return blocks

    # 文字层路径（原逻辑）
    for slide_no, texts in enumerate(pages, start=1):
        for t in texts:
            para += 1
            blocks.append(Block(page=slide_no, para_index=para, text=t))
    return blocks
