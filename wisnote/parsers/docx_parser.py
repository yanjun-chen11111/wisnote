"""Word 解析器：提取正文段落、表格、文本框与内嵌图片（OCR）

覆盖范围：
1. 正文段落 + 表格（python-docx 直接支持）；
2. 文本框文字（doc.paragraphs 不含文本框，从 body XML 提取，
   按 AlternateContent 双份结构去重）；
3. 内嵌图片（docx 本质是 zip，图片在包内；逐张 OCR，
   让"图里的字"也参与检索——流程图标注、截图说明等常见形态）。

不做：页眉页脚（多为页码/重复标题的噪声，重复内容会污染检索，
业界 RAG 惯例主动跳过）。

说明：.docx 格式本身不保存"页码"信息（分页由渲染决定），
因此 page 为 None，溯源时定位到"段落序号 + 原文片段"。
"""

import hashlib
from pathlib import Path

from docx import Document

from wisnote.parsers.base import Block

# 内嵌图 OCR 数量上限（防病态文件拖垮入库）
MAX_OCR_IMAGES = 100


def _textbox_texts(doc) -> list[str]:
    """文本框文字（w:txbxContent），按文本去重（mc:AlternateContent 会双份）"""
    seen: set[str] = set()
    texts: list[str] = []
    for node in doc.element.body.xpath(".//w:txbxContent//w:t"):
        t = (node.text or "").strip()
        if t and t not in seen:
            seen.add(t)
            texts.append(t)
    return texts


def _embedded_images(doc) -> list[bytes]:
    """内嵌图片二进制列表（按内容哈希去重——同一张图被引用多次只 OCR 一次）"""
    seen_hash: set[str] = set()
    blobs: list[bytes] = []
    for rel in doc.part.rels.values():
        if not rel.reltype.endswith("/image"):
            continue
        try:
            blob = rel.target_part.blob
        except Exception:
            continue  # 外部链接或损坏引用：跳过
        h = hashlib.sha1(blob).hexdigest()
        if h in seen_hash:
            continue
        seen_hash.add(h)
        blobs.append(blob)
    return blobs


def parse_docx(path: Path) -> list[Block]:
    blocks: list[Block] = []
    para = 0
    doc = Document(str(path))
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            para += 1
            blocks.append(Block(page=None, para_index=para, text=t))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                t = cell.text.strip()
                if t:
                    para += 1
                    blocks.append(Block(page=None, para_index=para, text=t))
    for t in _textbox_texts(doc):
        para += 1
        blocks.append(Block(page=None, para_index=para, text=t))

    # 内嵌图片 OCR（图里的字也参与检索）
    images = _embedded_images(doc)[:MAX_OCR_IMAGES]
    if images:
        from wisnote.parsers.image_parser import ocr_image_bytes

        for blob in images:
            for b in ocr_image_bytes(blob, page=None):
                para += 1
                blocks.append(Block(page=None, para_index=para, text=b.text))
    return blocks
