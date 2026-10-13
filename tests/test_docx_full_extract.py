"""docx 全量提取测试：内嵌图片 OCR + 文本框 + 正文回归"""

from docx import Document
from docx.oxml import parse_xml

from wisnote.parsers.docx_parser import parse_docx


def _make_text_image(path):
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("RGB", (800, 300), "white")
    d = ImageDraw.Draw(img)
    font = ImageFont.truetype(r"C:\Windows\Fonts\msyh.ttc", 40)
    d.text((40, 60), "内嵌图片里的字：银行家算法", fill="black", font=font)
    img.save(path)
    return path


def test_docx_embedded_image_ocr(tmp_path):
    """D1 含文字图片的 docx → 图片文字被 OCR 提取"""
    img = _make_text_image(tmp_path / "chart.png")
    doc = Document()
    doc.add_paragraph("正文段落：进程管理概述")
    doc.add_picture(str(img), width=__import__("docx").shared.Inches(6))
    p = tmp_path / "with_img.docx"
    doc.save(str(p))

    blocks = parse_docx(p)
    joined = " ".join(b.text for b in blocks)
    assert "正文段落" in joined, "正文不受影响"
    assert "银行家" in joined or "算法" in joined, f"内嵌图片文字应被 OCR 提取: {joined}"


def test_docx_textbox_extraction(tmp_path):
    """D2 文本框内容 → 从 XML 提取"""
    doc = Document()
    doc.add_paragraph("普通段落")
    # 注入一个 VML 文本框（python-docx 无 API，直接操作 XML）
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    V = "urn:schemas-microsoft-com:vml"
    xml = (
        f'<w:r xmlns:w="{W}" xmlns:v="{V}">'
        '<w:pict><v:shape style="width:100pt;height:50pt">'
        '<v:textbox><w:txbxContent><w:p><w:r><w:t>文本框独有内容XYZ789</w:t></w:r></w:p>'
        "</w:txbxContent></v:textbox></v:shape></w:pict></w:r>"
    )
    para = doc.add_paragraph()
    para._p.append(parse_xml(xml))
    p = tmp_path / "textbox.docx"
    doc.save(str(p))

    blocks = parse_docx(p)
    joined = " ".join(b.text for b in blocks)
    assert "XYZ789" in joined, f"文本框内容应被提取: {joined}"
    assert "普通段落" in joined


def test_docx_plain_regression(tmp_path):
    """D3 纯文字 docx → 原路径不受影响"""
    doc = Document()
    doc.add_paragraph("第一段内容")
    doc.add_paragraph("第二段内容")
    p = tmp_path / "plain.docx"
    doc.save(str(p))

    blocks = parse_docx(p)
    assert len(blocks) == 2
    assert all(b.page is None for b in blocks)
