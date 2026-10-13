"""OCR fallback 测试：扫描版 PDF / 图片型 PPT / 电子版不回归

覆盖场景（真实文件 + 真实 OCR，不用 mock——OCR 效果必须端到端验证）：
- E1 扫描版 PDF（文字渲染成图片嵌入 PDF，无电子层）→ OCR fallback 提取
- E2 电子版 PDF（文字层正常）→ 走原路径，不触发 OCR
- E3 图片型 PPT（整页一张图，无文本框）→ OCR fallback 提取
- E4 文字型 PPT → 走原路径，不触发 OCR
- E5 超页数上限的扫描 PDF → 明确报错（拒绝而非静默卡死）
"""

import pymupdf
import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from wisnote.parsers.pdf_parser import MAX_OCR_PAGES, parse_pdf
from wisnote.parsers.pptx_parser import MAX_OCR_SLIDES, parse_pptx


def _make_text_image(path, lines: list[tuple[int, str]]):
    """生成含打印体中文文字的 PNG"""
    from PIL import ImageDraw, ImageFont

    img = Image.new("RGB", (800, 400), "white")
    draw = ImageDraw.Draw(img)
    font = None
    for fp in (r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\simhei.ttf"):
        try:
            font = ImageFont.truetype(fp, 40)
            break
        except Exception:
            continue
    if font is None:
        font = ImageFont.load_default()
    for y, text in lines:
        draw.text((40, y), text, fill="black", font=font)
    img.save(path)
    return path


# ---------- 扫描版 PDF ----------

def test_scanned_pdf_ocr_fallback(tmp_path):
    """E1 扫描版 PDF：图片嵌入 PDF，电子层为空 → 渲染 OCR 提取成功"""
    img_path = _make_text_image(tmp_path / "scan.png", [(40, "磁盘调度算法SCAN"), (200, "最短寻道优先")])
    doc = pymupdf.open()
    page = doc.new_page(width=842, height=595)
    page.insert_image(page.rect, filename=str(img_path))
    pdf_path = tmp_path / "scanned.pdf"
    doc.save(str(pdf_path))
    doc.close()

    # 前置确认：确实是扫描版（电子层为空）
    with pymupdf.open(str(pdf_path)) as d:
        assert not d[0].get_text().strip(), "测试文件应是扫描版（无电子层）"

    blocks = parse_pdf(pdf_path)
    joined = " ".join(b.text for b in blocks)
    assert "磁盘" in joined, f"OCR fallback 应提取出文字，实际: {joined}"
    assert blocks[0].page == 1, "扫描版也应带页码"


def test_electronic_pdf_no_regression(tmp_path):
    """E2 电子版 PDF：文字层正常 → 原路径提取，含页码"""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 100), "electronic text layer works")
    pdf_path = tmp_path / "electronic.pdf"
    doc.save(str(pdf_path))
    doc.close()

    blocks = parse_pdf(pdf_path)
    assert any("electronic" in b.text for b in blocks)
    assert blocks[0].page == 1


def test_scanned_pdf_page_limit(tmp_path):
    """E5 超页数上限：明确 ValueError，而不是静默跑几十分钟"""
    img_path = _make_text_image(tmp_path / "scan.png", [(40, "占位文字")])
    doc = pymupdf.open()
    for _ in range(MAX_OCR_PAGES + 1):
        page = doc.new_page(width=842, height=595)
        page.insert_image(page.rect, filename=str(img_path))
    pdf_path = tmp_path / "too_big.pdf"
    doc.save(str(pdf_path))
    doc.close()

    with pytest.raises(ValueError, match="OCR 上限"):
        parse_pdf(pdf_path)


# ---------- 图片型 PPT ----------

def test_image_only_pptx_ocr_fallback(tmp_path):
    """E3 图片型 PPT：整页只有一张图片（无文本框）→ 渲染 OCR 提取"""
    img_path = _make_text_image(tmp_path / "slide.png", [(40, "进程调度优先级"), (200, "时间片轮转")])
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白版式
    slide.shapes.add_picture(str(img_path), Inches(0.5), Inches(0.5), Inches(9), Inches(5))
    ppt_path = tmp_path / "image_only.pptx"
    prs.save(str(ppt_path))

    # 前置确认：文字层为空
    prs2 = Presentation(str(ppt_path))
    for s in prs2.slides:
        for shape in s.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                pytest.fail("测试文件应是图片型（无文字层）")

    blocks = parse_pptx(ppt_path)
    joined = " ".join(b.text for b in blocks)
    assert "调度" in joined or "进程" in joined, f"OCR fallback 应提取出文字，实际: {joined}"


def test_text_pptx_no_regression(tmp_path):
    """E4 文字型 PPT：文本框正常 → 原路径提取"""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
    tb.text_frame.text = "正常文字型幻灯片内容"
    ppt_path = tmp_path / "text.pptx"
    prs.save(str(ppt_path))

    blocks = parse_pptx(ppt_path)
    assert any("文字型" in b.text for b in blocks)
    assert blocks[0].page == 1


def test_pptx_slide_limit_matches_pdf():
    """一致性：PPT 与 PDF 的 OCR 页数上限一致"""
    assert MAX_OCR_SLIDES == MAX_OCR_PAGES
